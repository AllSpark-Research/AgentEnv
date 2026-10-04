# -*- coding: utf-8 -*-
"""Deterministic verifier for the financial import-template task.

Run from the solver's final workspace: python3 verify.py
Prints exactly one JSON line and exits 0.
"""
import json
import os

RUBRIC_IDS = ["structure", "balance_sheet_fields", "profit_fields", "bank_account_fields"]
WEIGHTS = {"structure": 0.25, "balance_sheet_fields": 0.25, "profit_fields": 0.20, "bank_account_fields": 0.30}

HEADER_ROW = ["字段中文名", "字段英文名", "字段类型", "是否必填", "字段说明"]

# ---------------------------------------------------------------------------
# GROUND TRUTH (derived from 字段采集表.xlsx + 平台组通知.md + skill mapping
# library / type rules; see task_blueprint.json evaluation_plan)
# ---------------------------------------------------------------------------
EXPECTED = {
    "资产负债表_导入模板.xlsx": {
        "B1": "资产负债表",
        "B2": "balanceSheet",
        "rows": [
            ("编制单位", "reportingUnit", "单行文本", "是"),
            ("归属年月", "reportingPeriod", "单行文本", "是"),
            ("填报人", "reporter", "单行文本", "否"),
            ("填报日期", "reportDate", "日期时间", "是"),
            ("货币资金", "monetaryFunds", "金额", "是"),
            ("交易性金融资产", "tradingFinancialAssets", "金额", "否"),
            ("应收票据", "notesReceivable", "金额", "否"),
            ("应收账款", "accountsReceivable", "金额", "是"),
            ("预付款项", "prepayments", "金额", "否"),
            ("其他应收款", "otherReceivables", "金额", "否"),
            ("存货", "inventories", "金额", "是"),
            ("流动资产合计", "totalCurrentAssets", "金额", "是"),
            ("固定资产", "fixedAssets", "金额", "否"),
            ("无形资产", "intangibleAssets", "金额", "否"),
            ("资产总计", "totalAssets", "金额", "是"),
            ("短期借款", "shortTermBorrowings", "金额", "是"),
            ("应付票据", "notesPayable", "金额", "否"),
            ("应付账款", "accountsPayable", "金额", "是"),
            ("应付职工薪酬", "employeeBenefitsPayable", "金额", "否"),
            ("应交税费", "taxesPayable", "金额", "否"),
            ("其他应付款", "otherPayables", "金额", "否"),
            ("流动负债合计", "totalCurrentLiabilities", "金额", "是"),
            ("长期借款", "longTermBorrowings", "金额", "否"),
            ("负债合计", "totalLiabilities", "金额", "是"),
            ("实收资本（或股本）", "paidInCapital", "金额", "是"),
            ("资本公积", "capitalReserve", "金额", "否"),
            ("其他权益工具", "otherEquityInstruments", "金额", "否"),
            ("未分配利润", "retainedEarnings", "金额", "否"),
            ("所有者权益（或股东权益）合计", "totalOwnersEquity", "金额", "是"),
            ("负债和所有者权益（或股东权益）总计", "totalLiabilitiesAndEquity", "金额", "是"),
        ],
    },
    "利润表_导入模板.xlsx": {
        "B1": "利润表",
        "B2": "incomeStatement",
        "rows": [
            ("编制单位", "reportingUnit", "单行文本", "是"),
            ("归属年月", "reportingPeriod", "单行文本", "是"),
            ("营业收入", "operatingRevenue", "金额", "是"),
            ("减：营业成本", "operatingCosts", "金额", "是"),
            ("销售费用", "sellingExpenses", "金额", "否"),
            ("管理费用", "administrativeExpenses", "金额", "否"),
            ("研发费用", "rdExpenses", "金额", "否"),
            ("财务费用", "financialExpenses", "金额", "否"),
            ("加：其他收益", "otherIncome", "金额", "否"),
            ("投资收益", "investmentIncome", "金额", "否"),
            ("营业利润", "operatingProfit", "金额", "是"),
            ("利润总额", "totalProfit", "金额", "是"),
            ("减：所得税费用", "incomeTaxExpense", "金额", "是"),
            ("净利润", "netProfit", "金额", "是"),
        ],
    },
    "银行账号_导入模板.xlsx": {
        "B1": "银行账号",
        "B2": "bankAccount",
        "rows": [
            ("账号名称", "accountName", "单行文本", "是"),
            ("归属公司", "affiliatedCompany", "单行文本", "是"),
            ("账户类型", "accountType", "单选", "是"),
            ("账户状态", "accountStatus", "状态标签", "是"),
            ("银行名称", "bankName", "单行文本", "是"),
            ("开户行", "openingBank", "单行文本", "是"),
            ("银行账户", "bankAccount", "单行文本", "是"),
            ("币种", "currency", "单行文本", "是"),
            ("开户日期", "openingDate", "日期时间", "否"),
            ("账号用途", "accountUsage", "单行文本", "否"),
            ("账户限额", "field11", "单行文本", "是"),
        ],
    },
}

# Stale cash-flow leftovers kept in the blank template (rows 5-20); none of
# these Chinese names may appear in any deliverable.
STALE_LEFTOVER_NAMES = {
    "销售商品、提供劳务收到的现金", "收到的税费返还", "收到其他与经营活动有关的现金",
    "其中：财政补贴/政府项目回款", "经营活动现金流入小计", "购买商品、接受劳务支付的现金",
    "支付给职工以及为职工支付的现金", "支付的各项税费", "支付其他与经营活动有关的现金",
    "经营活动现金流出小计", "经营活动产生的现金流量净额", "期末现金及现金等价物余额",
    "汇总校验", "导入批次号", "期初现金及现金等价物余额", "汇率变动对现金的影响",
}
# Fields explicitly excluded by '不用新增' remarks in the collection workbook.
EXCLUDED_NAMES = {
    "其中：应收利息", "一年内到期的非流动资产", "一年内到期的非流动负债",
    "资产处置收益", "加：营业外收入", "营业外支出", "开户行地址", "开户行行号",
}


def cell_val(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        return None
    s = str(v).strip()
    return s if s != "" else None


def compare_table(filename, spec):
    """Return (field_score, structure_score, residue_ok, detail_dict)."""
    import openpyxl

    detail = {"file": filename}
    if not os.path.exists(filename):
        detail["error"] = "file missing"
        return 0.0, 0.0, False, detail
    try:
        wb = openpyxl.load_workbook(filename)
    except Exception as e:
        detail["error"] = f"unreadable xlsx: {e}"
        return 0.0, 0.0, False, detail
    ws = wb.active

    # --- structure checks ---
    st_checks = []
    st_checks.append(("A1 label", cell_val(ws, 1, 1) == "表中文名"))
    st_checks.append(("B1 table name", cell_val(ws, 1, 2) == spec["B1"]))
    st_checks.append(("A2 label", cell_val(ws, 2, 1) == "表英文名"))
    st_checks.append(("B2 entity name", cell_val(ws, 2, 2) == spec["B2"]))
    st_checks.append(("A3 label", cell_val(ws, 3, 1) == "说明"))
    hdr_ok = all(cell_val(ws, 4, c) == HEADER_ROW[c - 1] for c in range(1, 6))
    st_checks.append(("row4 headers", hdr_ok))
    structure_score = sum(1 for _, ok in st_checks if ok) / len(st_checks)
    detail["structure"] = {name: ok for name, ok in st_checks}

    # --- field rows ---
    exp = spec["rows"]
    got_n = len(exp)
    cell_hits, cell_total = 0, 0
    row_errors = []
    for i, (name, eng, ftype, req) in enumerate(exp):
        r = 5 + i
        got = [cell_val(ws, r, c) for c in range(1, 5)]
        want = [name, eng, ftype, req]
        for g, w in zip(got, want):
            cell_total += 1
            if g == w:
                cell_hits += 1
        if got != want:
            row_errors.append({"row": r, "got": got, "want": want})
    # extra rows beyond expected?
    extra_nonempty = []
    for r in range(5 + len(exp), max(ws.max_row + 1, 5 + len(exp) + 25)):
        for c in range(1, 6):
            v = cell_val(ws, r, c)
            if v is not None:
                extra_nonempty.append((r, c, v))
    # residue from stale leftover or excluded fields anywhere in col A?
    residue_hits = []
    for r in range(5, ws.max_row + 1):
        v = cell_val(ws, r, 1)
        if v in STALE_LEFTOVER_NAMES or v in EXCLUDED_NAMES:
            residue_hits.append((r, v))
    residue_ok = (len(extra_nonempty) == 0) and (len(residue_hits) == 0)

    field_score = (cell_hits / cell_total) if cell_total else 0.0
    if extra_nonempty:
        field_score *= 0.5  # leaked stale/extra rows are a serious defect
    detail["cell_hits"] = f"{cell_hits}/{cell_total}"
    detail["row_errors_sample"] = row_errors[:5]
    detail["extra_nonempty_sample"] = extra_nonempty[:5]
    detail["residue_hits"] = residue_hits[:5]
    detail["rows_expected"] = got_n
    return field_score, structure_score, residue_ok, detail


def main():
    import openpyxl  # noqa: F401 -- availability probe

    results = {}
    details = {}
    field_ids = {
        "资产负债表_导入模板.xlsx": "balance_sheet_fields",
        "利润表_导入模板.xlsx": "profit_fields",
        "银行账号_导入模板.xlsx": "bank_account_fields",
    }
    structure_parts = []
    residue_fails = []

    for fname, spec in EXPECTED.items():
        try:
            fs, ss, rok, det = compare_table(fname, spec)
        except Exception as e:
            fs, ss, rok, det = 0.0, 0.0, False, {"file": fname, "error": str(e)}
        results[field_ids[fname]] = fs
        details[field_ids[fname]] = det
        structure_parts.append(ss)
        if not rok:
            residue_fails.append(fname)

    # structure item: mean of per-file structures AND gate on residue
    struct_score = sum(structure_parts) / len(structure_parts) if structure_parts else 0.0
    residue_pen = 1.0 - 0.5 * (len(residue_fails) / max(len(EXPECTED), 1))
    struct_total = max(0.0, struct_score * residue_pen)
    details["structure"] = {
        "per_file_structure": structure_parts,
        "residue_fail_files": residue_fails,
        "note": "mean structure score penalized by files containing stale/excluded residue or extra rows",
    }

    per_item = {
        "structure": struct_total,
        "balance_sheet_fields": results["balance_sheet_fields"],
        "profit_fields": results["profit_fields"],
        "bank_account_fields": results["bank_account_fields"],
    }
    checks = []
    for rid in RUBRIC_IDS:
        checks.append({
            "id": rid,
            "weight": WEIGHTS[rid],
            "score": per_item[rid],
            "detail": json.dumps(details.get(rid, {}), ensure_ascii=False)[:600],
        })
    total_w = sum(WEIGHTS[rid] for rid in RUBRIC_IDS)
    score = sum(WEIGHTS[rid] * per_item[rid] for rid in RUBRIC_IDS) / total_w

    print(json.dumps({"per_item": per_item, "checks": checks, "score": score}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        # never crash: report zeros for everything
        per_item = {rid: 0.0 for rid in RUBRIC_IDS}
        checks = [{"id": rid, "weight": WEIGHTS[rid], "score": 0.0, "detail": f"verifier error: {e}"} for rid in RUBRIC_IDS]
        print(json.dumps({"per_item": per_item, "checks": checks, "score": 0.0}, ensure_ascii=False))
